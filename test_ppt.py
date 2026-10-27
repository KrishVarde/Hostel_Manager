import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_layout = prs.slide_layouts[6] # completely blank
slide = prs.slides.add_slide(blank_layout)

# Test title slide
# Dark navy banner
banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
banner.fill.solid()
banner.fill.fore_color.rgb = RGBColor(27, 54, 93) # #1B365D
banner.line.color.rgb = RGBColor(27, 54, 93)

print("Title slide created successfully")
prs.save("test_presentation.pptx")
print("Saved test_presentation.pptx")
